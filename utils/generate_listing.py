import argparse
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


CODE_FONT = "Cascadia Mono"
CODE_SIZE_PT = 9.5
HEADER_PRIORITY = {
    "Person.h": 0,
    "Employee.h": 1,
    "Worker.h": 2,
}
HEADER_EXTENSIONS = {".h", ".hh", ".hpp", ".hxx"}
SOURCE_EXTENSIONS = {".c", ".cc", ".cpp", ".cxx"}


def set_run_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold


def add_code_block(doc: Document, title: str, code: str) -> None:
    title_par = doc.add_paragraph()
    title_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_par.add_run(title)
    set_run_font(title_run, CODE_FONT, CODE_SIZE_PT, bold=True)

    code_par = doc.add_paragraph()
    code_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code_par.paragraph_format.space_before = Pt(0)
    code_par.paragraph_format.space_after = Pt(0)
    lines = code.splitlines()
    for idx, line in enumerate(lines):
        code_run = code_par.add_run(line)
        set_run_font(code_run, CODE_FONT, CODE_SIZE_PT, bold=False)
        if idx != len(lines) - 1:
            code_run.add_break()


def collect_files(project_root: Path, bases: Iterable[Path]) -> list[Path]:
    headers: list[Path] = []
    sources: list[Path] = []

    for base in bases:
        if not base.exists():
            continue
        for file_path in base.rglob("*"):
            if file_path.suffix in HEADER_EXTENSIONS:
                headers.append(file_path)
            elif file_path.suffix in SOURCE_EXTENSIONS:
                sources.append(file_path)

    def header_sort_key(path: Path) -> tuple[int, str]:
        priority = HEADER_PRIORITY.get(path.name, len(HEADER_PRIORITY))
        relative = path.relative_to(project_root).as_posix()
        return priority, relative

    headers_sorted = sorted(headers, key=header_sort_key)
    sources_sorted = sorted(
        sources, key=lambda p: p.relative_to(project_root).as_posix()
    )
    return headers_sorted + sources_sorted


def build_listing(output_path: Path, project_root: Path, folders: Sequence[str]) -> None:
    doc = Document()
    files = collect_files(project_root, [project_root / folder for folder in folders])
    for file_path in files:
        try:
            text = file_path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        relative_name = file_path.relative_to(project_root).as_posix()
        add_code_block(doc, relative_name, text)
        doc.add_paragraph(" ")

    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    default_root = Path(__file__).resolve().parents[1]
    default_output = default_root / "utils" / "Листинг_кода.docx"
    parser = argparse.ArgumentParser(
        description="Собрать листинг проекта в docx с кодовым шрифтом."
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=default_output,
        help="Путь до docx-файла (по умолчанию utils/Листинг_кода.docx)",
    )
    parser.add_argument(
        "-r",
        "--root",
        type=Path,
        default=default_root,
        help="Корень проекта (по умолчанию каталог репозитория)",
    )
    parser.add_argument(
        "-d",
        "--dirs",
        nargs="+",
        default=["include", "src"],
        help="Каталоги, из которых собирать код (по умолчанию include и src)",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_listing(args.output, args.root, args.dirs)
    print(f"Готово: {args.output}")
